#!/usr/bin/env python3
"""Refillable Quest / IMU / Neon status workbook.

Writes an Excel book (and CSV copies) you can regenerate any time:

  usable                 people with all 6 walking cells Quest+IMU+Neon TRUE; duplicate if another usable person has the same study order
  walking                one row per person, TRUE/FALSE for each walking cell × stream
  bouts                  long table (all 12 folders) for filtering
  openeye_eval           one row per bout: OpenEye eval median/mean (practice inherited for walking)
  study_order            one row per person; fill in session order (kept on regenerate)
  trial_counts_by_bout   Fitts trial counts from check_mt_dwell (participant × bout × modality)
  trial_counts_summary   totals: all / per person / modality / layout / standing·walking

Usage (from scripts/02_analysis/):
    uv run python participant_status.py
    uv run python participant_status.py --first 23 --last 56
"""
from __future__ import annotations

from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _bootstrap import ensure_analysis_path

ensure_analysis_path()

import argparse

import pandas as pd

from _paths import (
    BOUTS,
    DATA_ROOT,
    INTERACTIONS,
    PARTICIPANTS,
    STAGE_DIRS,
    WALKING_BOUTS,
    analysis_out,
    bout_dir,
    bout_labels,
    is_practice_bout,
    participant_dir,
)
from bout_quality import find_blinks_csv, load_eval_deg, load_grid_fracs, load_sync_info
from fitts_gait_onset import pick_quest_json
from gait_onset import GAIT_REF_FOOT_OVERRIDE

EVAL_MAX_DEG = 3.0
MIN_ICS = 40
OUT_XLSX = DATA_ROOT / "participants" / "participant_status.xlsx"
OUT_DIR = analysis_out(__file__)

INTER_SHORT = {"HeadPinch": "Head", "HandPinch": "Hand", "EyePinch": "Eye"}
BOUT_SHORT = {
    "Ring": "Ring",
    "Rectangle": "Rect",
    "PracticeRing": "PRing",
    "PracticeRectangle": "PRect",
}

STUDY_ORDER_SHEET = "study_order"
EVAL_SHEET = "openeye_eval"
TRIAL_BY_BOUT_SHEET = "trial_counts_by_bout"
TRIAL_SUMMARY_SHEET = "trial_counts_summary"
MT_DWELL_SUMMARY = analysis_out("02_05_cursor_stability/check_mt_dwell.py") / "summary.csv"
ORDER_ID_COL = "participant"
ORDER_FILL_COLS = (
    "sex",
    "block_1",
    "block_2",
    "layout_1",
    "layout_2",
    "modality_1",
    "modality_2",
    "modality_3",
    "notes",
)
ORDER_SEX_COLS = ("sex",)
ORDER_BLOCK_COLS = ("block_1", "block_2")
ORDER_LAYOUT_COLS = ("layout_1", "layout_2")
ORDER_MODALITY_COLS = ("modality_1", "modality_2", "modality_3")
ORDER_KEY_COLS = ORDER_BLOCK_COLS + ORDER_LAYOUT_COLS + ORDER_MODALITY_COLS
DEPRECATED_ORDER_COLS = {
    "male",
    "female",
    "walk_first_layout",
    "stand_first_layout",
    "Ring_1",
    "Ring_2",
    "Ring_3",
    "Rect_1",
    "Rect_2",
    "Rect_3",
    "PRing_1",
    "PRing_2",
    "PRing_3",
    "PRect_1",
    "PRect_2",
    "PRect_3",
}


def _n_ics_ok(path: Path) -> int:
    if not path.is_file():
        return 0
    df = pd.read_csv(path)
    if df.empty:
        return 0
    if "is_outlier" in df.columns:
        df = df[df["is_outlier"] == False]  # noqa: E712
    return int(len(df))


def _quest_ok(bout: Path) -> bool:
    try:
        pick_quest_json(bout)
        return True
    except (FileNotFoundError, FileExistsError, OSError):
        return False


def _imu_ok(bout: Path) -> bool:
    """Both feet recorded and usable. Practice = no foot IMU → False."""
    speed = bout.parent.name
    if is_practice_bout(speed):
        return False
    subject, run = bout_labels(bout)
    if (subject, run) in GAIT_REF_FOOT_OVERRIDE:
        return False
    proc = bout / STAGE_DIRS["gait"] / "processed" / subject / run
    if (proc / "RF_DEAD_STUB.txt").is_file():
        return False
    xsens = bout / STAGE_DIRS["gait_xsens"]
    lf_x = (xsens / "LF.csv").is_file()
    rf_x = (xsens / "RF.csv").is_file()
    if not (lf_x and rf_x):
        return False
    lf_p = proc / "left_foot_core_params.csv"
    rf_p = proc / "right_foot_core_params.csv"
    if lf_p.is_file() or rf_p.is_file():
        n_lf = _n_ics_ok(lf_p)
        n_rf = _n_ics_ok(rf_p)
        if n_lf < 1 or n_rf < 1:
            return False
        if n_lf < MIN_ICS or n_rf < MIN_ICS:
            return False
        if n_rf and (n_lf / n_rf < 0.5 or n_lf / n_rf > 2.0):
            return False
    return True


def _neon_ok(bout: Path) -> bool:
    if find_blinks_csv(bout) is not None:
        return True
    fr = load_grid_fracs(bout)
    g = fr.get("gaze_valid_frac")
    if pd.notna(g) and float(g) >= 0.99:
        return True
    for stage in ("corrected", "cleaned", "grid"):
        if (bout / STAGE_DIRS[stage] / "gaze.csv").is_file():
            return True
    return False


def _list_ids(first: int, last: int) -> list[int]:
    out: list[int] = []
    for n in range(int(first), int(last) + 1):
        if participant_dir(n).is_dir():
            out.append(n)
    return out


def scan_bouts(ids: list[int]) -> pd.DataFrame:
    rows: list[dict] = []
    for n in ids:
        pid = f"participant{n}"
        for layout in BOUTS:
            for inter in INTERACTIONS:
                bout = bout_dir(n, layout, inter)
                standing = is_practice_bout(layout)
                if not bout.is_dir():
                    rows.append(
                        {
                            "participant": n,
                            "id": pid,
                            "bout": layout,
                            "interaction": inter,
                            "standing": standing,
                            "quest": False,
                            "imu": False,
                            "neon": False,
                            "sync": False,
                            "eval_ok": False if inter == "EyePinch" else True,
                            "eval_median_deg": float("nan"),
                            "note": "no folder",
                        }
                    )
                    continue
                quest = _quest_ok(bout)
                imu = _imu_ok(bout)
                neon = _neon_ok(bout)
                sync = bool(load_sync_info(bout).get("has_offset_quest_to_pc"))
                ev = load_eval_deg(bout)
                med = ev.get("eval_median_deg")
                if inter == "EyePinch":
                    eval_ok = bool(pd.notna(med) and float(med) < EVAL_MAX_DEG)
                else:
                    eval_ok = True
                notes: list[str] = []
                if standing:
                    notes.append("standing: IMU unused")
                if not quest:
                    notes.append("no Quest JSON")
                if not imu and not standing:
                    notes.append("IMU fail")
                if not neon:
                    notes.append("no Neon")
                if inter == "EyePinch" and not eval_ok:
                    notes.append(f"eval_median={med}")
                rows.append(
                    {
                        "participant": n,
                        "id": pid,
                        "bout": layout,
                        "interaction": inter,
                        "standing": standing,
                        "quest": bool(quest),
                        "imu": bool(imu),
                        "neon": bool(neon),
                        "sync": bool(sync),
                        "eval_ok": bool(eval_ok),
                        "eval_median_deg": med if pd.notna(med) else float("nan"),
                        "note": "; ".join(notes),
                    }
                )
    return pd.DataFrame(rows)


def scan_eval(ids: list[int]) -> pd.DataFrame:
    """One row per bout folder: OpenEye ridge eval (median gate 3°)."""
    rows: list[dict] = []
    for n in ids:
        for layout in BOUTS:
            for inter in INTERACTIONS:
                bout = bout_dir(n, layout, inter)
                standing = is_practice_bout(layout)
                if not bout.is_dir():
                    rows.append(
                        {
                            "participant": n,
                            "bout": layout,
                            "interaction": INTER_SHORT[inter],
                            "standing": standing,
                            "source": "no_folder",
                            "median_deg": float("nan"),
                            "mean_deg": float("nan"),
                            "std_deg": float("nan"),
                            "n": float("nan"),
                            "eval_ok": False,
                            "datetime": "",
                        }
                    )
                    continue
                ev = load_eval_deg(bout)
                med = ev.get("eval_median_deg")
                ok = bool(pd.notna(med) and float(med) < EVAL_MAX_DEG)
                n_samp = ev.get("eval_n")
                rows.append(
                    {
                        "participant": n,
                        "bout": layout,
                        "interaction": INTER_SHORT[inter],
                        "standing": standing,
                        "source": ev.get("eval_source") or "missing",
                        "median_deg": med if pd.notna(med) else float("nan"),
                        "mean_deg": ev.get("eval_mean_deg", float("nan")),
                        "std_deg": ev.get("eval_std_deg", float("nan")),
                        "n": n_samp if pd.notna(n_samp) else float("nan"),
                        "eval_ok": ok,
                        "datetime": ev.get("eval_datetime") or "",
                    }
                )
    return pd.DataFrame(rows)


def walking_wide(bouts: pd.DataFrame) -> pd.DataFrame:
    w = bouts[~bouts["standing"]].copy()
    rows = []
    for n, g in w.groupby("participant"):
        rec: dict = {"participant": int(n)}
        n_ok = 0
        for layout in WALKING_BOUTS:
            for inter in INTERACTIONS:
                prefix = f"{BOUT_SHORT[layout]}_{INTER_SHORT[inter]}"
                hit = g[(g["bout"] == layout) & (g["interaction"] == inter)]
                if hit.empty:
                    rec[f"{prefix}_quest"] = False
                    rec[f"{prefix}_imu"] = False
                    rec[f"{prefix}_neon"] = False
                    continue
                q, i, ne = bool(hit["quest"].iloc[0]), bool(hit["imu"].iloc[0]), bool(hit["neon"].iloc[0])
                rec[f"{prefix}_quest"] = q
                rec[f"{prefix}_imu"] = i
                rec[f"{prefix}_neon"] = ne
                if q and i and ne:
                    n_ok += 1
        rec["n_cells_ok"] = n_ok
        rec["usable"] = n_ok == 6
        rows.append(rec)
    return pd.DataFrame(rows)


SKIP_RED = {
    "participant",
    "id",
    "bout",
    "interaction",
    "standing",
    "note",
    "eval_median_deg",
    "n_cells_ok",
    "item",
    "meaning",
    "duplicate",
    "same_order_as",
}


def _filled(val) -> bool:
    if val is None:
        return False
    try:
        if pd.isna(val):
            return False
    except (TypeError, ValueError):
        pass
    return str(val).strip() != ""


def _norm_order_val(val) -> str:
    return str(val).strip() if _filled(val) else ""


def usable_order_flags(usable: pd.DataFrame, order: pd.DataFrame) -> pd.DataFrame:
    """Flag usable people who share block / layout / modality order with another usable person."""
    out = usable[["participant"]].copy()
    out["duplicate"] = False
    out["same_order_as"] = pd.NA
    if order is None or order.empty or "participant" not in order.columns:
        return out
    o = order.copy()
    o["participant"] = pd.to_numeric(o["participant"], errors="coerce")
    o = o.dropna(subset=["participant"])
    o["participant"] = o["participant"].astype(int)
    keep = set(int(n) for n in out["participant"])
    o = o[o["participant"].isin(keep)].copy()
    for c in ORDER_KEY_COLS:
        if c not in o.columns:
            o[c] = ""
        o[c] = o[c].map(_norm_order_val)
    complete = o[o[list(ORDER_KEY_COLS)].ne("").all(axis=1)]
    if complete.empty:
        return out
    complete = complete.copy()
    complete["_key"] = complete[list(ORDER_KEY_COLS)].agg("|".join, axis=1)
    dup_map: dict[int, str] = {}
    for _, series in complete.groupby("_key")["participant"]:
        ids = sorted(int(x) for x in series.tolist())
        if len(ids) < 2:
            continue
        for pid in ids:
            dup_map[pid] = ", ".join(f"p{x}" for x in ids if x != pid)
    out["duplicate"] = [int(n) in dup_map for n in out["participant"]]
    out["same_order_as"] = [dup_map.get(int(n), pd.NA) for n in out["participant"]]
    return out


def empty_study_order(ids: list[int]) -> pd.DataFrame:
    recs = []
    for n in ids:
        rec = {ORDER_ID_COL: int(n)}
        rec.update({c: pd.NA for c in ORDER_FILL_COLS})
        recs.append(rec)
    return pd.DataFrame(recs)


def load_study_order(path: Path) -> pd.DataFrame:
    if not path.is_file():
        return pd.DataFrame()
    try:
        df = pd.read_excel(path, sheet_name=STUDY_ORDER_SHEET)
    except (ValueError, KeyError, OSError):
        return pd.DataFrame()
    if df is None or df.empty or ORDER_ID_COL not in df.columns:
        return pd.DataFrame()
    return df


def merge_study_order(ids: list[int], saved: pd.DataFrame) -> pd.DataFrame:
    """Keep typed cells; add new people as blank rows. Extra columns are kept."""
    base = empty_study_order(ids).set_index(ORDER_ID_COL)
    if saved is None or saved.empty or ORDER_ID_COL not in saved.columns:
        return base.reset_index()
    saved = saved.copy()
    saved[ORDER_ID_COL] = pd.to_numeric(saved[ORDER_ID_COL], errors="coerce")
    saved = saved.dropna(subset=[ORDER_ID_COL])
    saved[ORDER_ID_COL] = saved[ORDER_ID_COL].astype(int)
    saved = saved.drop_duplicates(ORDER_ID_COL, keep="last").set_index(ORDER_ID_COL)
    extra_ids = [i for i in saved.index if i not in base.index]
    if extra_ids:
        extra = pd.DataFrame(index=extra_ids)
        extra.index.name = ORDER_ID_COL
        for c in ORDER_FILL_COLS:
            extra[c] = pd.NA
        base = pd.concat([base, extra])
    for col in saved.columns:
        if col not in base.columns:
            base[col] = pd.NA
        for pid, val in saved[col].items():
            if _filled(val):
                base.at[pid, col] = val
    ordered = [i for i in ids if i in base.index] + extra_ids
    keep_cols = [c for c in ORDER_FILL_COLS if c in base.columns]
    extra_cols = [
        c
        for c in base.columns
        if c not in keep_cols
        and c not in DEPRECATED_ORDER_COLS
        and any(_filled(v) for v in base[c])
    ]
    return base.loc[ordered, keep_cols + extra_cols].reset_index()


def _header_cols(ws) -> dict[str, int]:
    return {str(ws.cell(1, col).value): col for col in range(1, ws.max_column + 1) if ws.cell(1, col).value}


def _paint_order_blanks(ws) -> None:
    from openpyxl.styles import PatternFill

    fill = PatternFill("solid", fgColor="FFF2CC")
    headers = _header_cols(ws)
    skip = {ORDER_ID_COL}
    for row in range(2, ws.max_row + 1):
        for name, col in headers.items():
            if name in skip:
                continue
            cell = ws.cell(row, col)
            if not _filled(cell.value):
                cell.fill = fill


def _order_dropdowns(ws) -> None:
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    headers = _header_cols(ws)
    last = ws.max_row if ws.max_row >= 2 else 2

    def add_list(colnames: tuple[str, ...], choices: str) -> None:
        cells = []
        for name in colnames:
            col = headers.get(name)
            if not col:
                continue
            letter = get_column_letter(col)
            cells.append(f"{letter}2:{letter}{last}")
        if not cells:
            return
        dv = DataValidation(type="list", formula1=choices, allow_blank=True)
        dv.error = "Pick from the list"
        dv.errorTitle = "Invalid value"
        dv.prompt = "Pick from the list (or leave blank)"
        dv.promptTitle = "Study order"
        for rng in cells:
            dv.add(rng)
        ws.add_data_validation(dv)

    add_list(ORDER_SEX_COLS, '"Male,Female"')
    add_list(ORDER_BLOCK_COLS, '"Practice,Main"')
    add_list(ORDER_LAYOUT_COLS, '"Ring,Rect"')
    add_list(ORDER_MODALITY_COLS, '"Head,Hand,Eye"')


def _paint_duplicate_true(ws) -> None:
    from openpyxl.styles import Font, PatternFill

    fill = PatternFill("solid", fgColor="FCE4D6")
    font = Font(color="9C5700")
    headers = _header_cols(ws)
    col = headers.get("duplicate")
    if not col:
        return
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row, col)
        if cell.value is True:
            cell.fill = fill
            cell.font = font


def _paint_false(ws) -> None:
    from openpyxl.styles import Font, PatternFill

    fill = PatternFill("solid", fgColor="FFC7CE")
    font = Font(color="9C0006")
    headers = {
        col: ws.cell(1, col).value
        for col in range(1, ws.max_column + 1)
    }
    for row in range(2, ws.max_row + 1):
        for col, header in headers.items():
            if header in SKIP_RED:
                continue
            cell = ws.cell(row, col)
            if cell.value is False:
                cell.fill = fill
                cell.font = font


def _paint_eval(ws) -> None:
    from openpyxl.styles import Font, PatternFill

    bad = PatternFill("solid", fgColor="FFC7CE")
    bad_font = Font(color="9C0006")
    miss = PatternFill("solid", fgColor="FFF2CC")
    headers = _header_cols(ws)
    src_col = headers.get("source")
    ok_col = headers.get("eval_ok")
    med_col = headers.get("median_deg")
    for row in range(2, ws.max_row + 1):
        if src_col:
            src = ws.cell(row, src_col)
            if str(src.value or "") in {"missing", "no_folder"}:
                src.fill = miss
        ok = ws.cell(row, ok_col).value is True if ok_col else False
        if ok_col and ws.cell(row, ok_col).value is False:
            ws.cell(row, ok_col).fill = bad
            ws.cell(row, ok_col).font = bad_font
        if med_col and not ok:
            cell = ws.cell(row, med_col)
            cell.fill = bad
            cell.font = bad_font


def _pid_int(label: object) -> int | None:
    s = str(label).strip().removeprefix("participant").removeprefix("p")
    return int(s) if s.isdigit() else None


def load_trial_counts(ids: list[int]) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Fitts trial counts from check_mt_dwell summary (selections + successes).

    Returns (by_bout, summary). Empty if summary.csv is missing.
    """
    if not MT_DWELL_SUMMARY.is_file():
        return pd.DataFrame(), pd.DataFrame()
    raw = pd.read_csv(MT_DWELL_SUMMARY)
    if "participant" not in raw.columns:
        return pd.DataFrame(), pd.DataFrame()
    want = set(ids)
    raw = raw.copy()
    raw["participant"] = raw["participant"].map(_pid_int)
    df = raw[raw["participant"].isin(want)].copy()
    if df.empty:
        return pd.DataFrame(), pd.DataFrame()
    keep_cols = [
        c
        for c in (
            "participant",
            "speed",
            "interaction",
            "layout",
            "speed_group",
            "n_iso_selections",
            "n_success",
            "hit_rate",
            "n_selections",
            "n_with_first_hit",
        )
        if c in df.columns
    ]
    sort_cols = [c for c in ("participant", "speed_group", "layout", "interaction") if c in keep_cols]
    by_bout = df[keep_cols].sort_values(sort_cols).reset_index(drop=True)

    rows: list[dict] = []
    n_sel = pd.to_numeric(by_bout["n_iso_selections"], errors="coerce").fillna(0)
    n_ok = pd.to_numeric(by_bout["n_success"], errors="coerce").fillna(0)
    rows.append(
        {
            "level": "total",
            "key": "all",
            "n_iso_selections": int(n_sel.sum()),
            "n_success": int(n_ok.sum()),
            "n_bouts": int(len(by_bout)),
        }
    )
    for level, col in (
        ("participant", "participant"),
        ("modality", "interaction"),
        ("layout", "layout"),
        ("standing_walking", "speed_group"),
    ):
        if col not in by_bout.columns:
            continue
        for key, g in by_bout.groupby(col, dropna=False):
            rows.append(
                {
                    "level": level,
                    "key": str(key),
                    "n_iso_selections": int(
                        pd.to_numeric(g["n_iso_selections"], errors="coerce").fillna(0).sum()
                    ),
                    "n_success": int(
                        pd.to_numeric(g["n_success"], errors="coerce").fillna(0).sum()
                    ),
                    "n_bouts": int(len(g)),
                }
            )
    if {"participant", "speed_group"} <= set(by_bout.columns):
        for (pid, sg), g in by_bout.groupby(["participant", "speed_group"], dropna=False):
            rows.append(
                {
                    "level": "participant_x_standing_walking",
                    "key": f"{pid}|{sg}",
                    "n_iso_selections": int(
                        pd.to_numeric(g["n_iso_selections"], errors="coerce").fillna(0).sum()
                    ),
                    "n_success": int(
                        pd.to_numeric(g["n_success"], errors="coerce").fillna(0).sum()
                    ),
                    "n_bouts": int(len(g)),
                }
            )
    return by_bout, pd.DataFrame(rows)


def write_xlsx(
    usable: pd.DataFrame,
    walking: pd.DataFrame,
    bouts: pd.DataFrame,
    order: pd.DataFrame,
    evals: pd.DataFrame,
    path: Path,
    *,
    trial_by_bout: pd.DataFrame | None = None,
    trial_summary: pd.DataFrame | None = None,
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    legend = pd.DataFrame(
        {
            "item": [
                "quest",
                "imu",
                "neon",
                "usable",
                "usable duplicate",
                "standing imu",
                "FALSE cells",
                "openeye_eval",
                "study_order",
                "study_order sex",
                "study_order yellow",
                "trial_counts_by_bout",
                "trial_counts_summary",
                "regenerate",
            ],
            "meaning": [
                "TRUE if the matching Quest Fitts JSON is present (one file).",
                "TRUE if both feet are usable (Xsens LF+RF, no RF dead-stub / single-foot override; if gait ran, ≥40 ICs/foot and LF/RF ratio 0.5–2).",
                "TRUE if blinks.csv exists, or grid Neon gaze frac ≥ 0.99, or gaze.csv after export.",
                "TRUE on the usable sheet if all 6 walking cells are Quest AND IMU AND Neon TRUE.",
                "TRUE if another usable person has the same Practice/Main × Ring/Rect × Head/Hand/Eye order. same_order_as lists those IDs. Sex is ignored.",
                "Always FALSE (no foot IMU by design).",
                "Light red fill on walking / bouts sheets.",
                "One row per bout. median_deg from ridge_biquadratic. eval_ok = median < 3°. Walking usually source=practice (same interaction). Red = fail or missing.",
                "Fill-in sheet, nested order: Practice vs Main (block_1/2), then Ring vs Rect (layout_1/2), then Head/Hand/Eye (modality_1/2/3).",
                "Physical sex: Male or Female dropdown.",
                "Yellow cells are empty — type or use the dropdown. Filled cells are kept when you regenerate this workbook.",
                "From check_mt_dwell summary.csv: one row per participant × bout × modality. n_iso_selections = Fitts trials; n_success = successful confirms.",
                "Aggregated trial counts: total, per participant, modality, layout, standing/walking, and participant×standing/walking.",
                "From scripts/02_analysis:  uv run python participant_status.py --first 23 --last 56",
            ],
        }
    )
    tmp = path.with_suffix(".tmp.xlsx")
    with pd.ExcelWriter(tmp, engine="openpyxl") as writer:
        usable.to_excel(writer, sheet_name="usable", index=False)
        walking.to_excel(writer, sheet_name="walking", index=False)
        bouts.to_excel(writer, sheet_name="bouts", index=False)
        evals.to_excel(writer, sheet_name=EVAL_SHEET, index=False)
        order.to_excel(writer, sheet_name=STUDY_ORDER_SHEET, index=False)
        if trial_by_bout is not None and not trial_by_bout.empty:
            trial_by_bout.to_excel(writer, sheet_name=TRIAL_BY_BOUT_SHEET, index=False)
        if trial_summary is not None and not trial_summary.empty:
            trial_summary.to_excel(writer, sheet_name=TRIAL_SUMMARY_SHEET, index=False)
        legend.to_excel(writer, sheet_name="legend", index=False)
        wb = writer.book
        sheet_names = [
            "usable",
            "walking",
            "bouts",
            EVAL_SHEET,
            STUDY_ORDER_SHEET,
            TRIAL_BY_BOUT_SHEET,
            TRIAL_SUMMARY_SHEET,
            "legend",
        ]
        for name in sheet_names:
            if name not in wb.sheetnames:
                continue
            ws = wb[name]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
        _paint_false(wb["walking"])
        _paint_false(wb["bouts"])
        _paint_eval(wb[EVAL_SHEET])
        _paint_duplicate_true(wb["usable"])
        wb["usable"].column_dimensions["C"].width = 16
        order_ws = wb[STUDY_ORDER_SHEET]
        _paint_order_blanks(order_ws)
        _order_dropdowns(order_ws)
        from openpyxl.utils import get_column_letter

        order_ws.column_dimensions["A"].width = 12
        for col in range(2, order_ws.max_column):
            order_ws.column_dimensions[get_column_letter(col)].width = 18
        order_ws.column_dimensions[get_column_letter(order_ws.max_column)].width = 32
    try:
        tmp.replace(path)
    except OSError:
        fallback = path.with_name(path.stem + "_new.xlsx")
        tmp.replace(fallback)
        print(f"Could not overwrite {path.name} (file open?). Wrote {fallback.name} instead.")
        return fallback
    return path


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--first", type=int, default=23)
    p.add_argument("--last", type=int, default=56)
    p.add_argument("--out", type=Path, default=OUT_XLSX)
    return p.parse_args()


def main() -> None:
    args = parse_args()
    ids = _list_ids(args.first, args.last)
    if not ids:
        raise SystemExit(f"no participant folders in {args.first}–{args.last}")
    bouts = scan_bouts(ids)
    walking = walking_wide(bouts)
    usable = walking.loc[walking["usable"], ["participant"]].reset_index(drop=True)
    order = merge_study_order(ids, load_study_order(args.out))
    usable = usable_order_flags(usable, order)
    evals = scan_eval(ids)
    trial_by_bout, trial_summary = load_trial_counts(ids)

    out_xlsx = args.out
    written = write_xlsx(
        usable,
        walking,
        bouts,
        order,
        evals,
        out_xlsx,
        trial_by_bout=trial_by_bout,
        trial_summary=trial_summary,
    )
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    usable.to_csv(OUT_DIR / "usable_participants.csv", index=False)
    walking.to_csv(OUT_DIR / "status_walking.csv", index=False)
    bouts.to_csv(OUT_DIR / "status_bouts.csv", index=False)
    order.to_csv(OUT_DIR / "study_order.csv", index=False)
    evals.to_csv(OUT_DIR / "openeye_eval.csv", index=False)
    if not trial_by_bout.empty:
        trial_by_bout.to_csv(OUT_DIR / "trial_counts_by_bout.csv", index=False)
    if not trial_summary.empty:
        trial_summary.to_csv(OUT_DIR / "trial_counts_summary.csv", index=False)

    print(f"Usable (6/6 Quest+IMU+Neon): {len(usable)}  -> {usable['participant'].tolist()}")
    dups = usable.loc[usable["duplicate"], ["participant", "same_order_as"]]
    if not dups.empty:
        pairs = ", ".join(f"p{int(r.participant)}={r.same_order_as}" for r in dups.itertuples())
        print(f"Duplicate study order among usable: {pairs}")
    if trial_summary.empty:
        print(f"Trial counts: skipped (missing {MT_DWELL_SUMMARY})")
    else:
        tot = trial_summary.loc[trial_summary["level"] == "total"].iloc[0]
        print(
            f"Trial counts: {int(tot['n_iso_selections'])} selections / "
            f"{int(tot['n_success'])} success across {int(tot['n_bouts'])} bouts "
            f"({trial_by_bout['participant'].nunique()} people in check_mt_dwell)"
        )
    print(f"Wrote {written}")
    print(f"CSV copies under {OUT_DIR}")


if __name__ == "__main__":
    main()
